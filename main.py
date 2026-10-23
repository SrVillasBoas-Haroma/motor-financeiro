"""
Motor Financeiro - Haroma
Calcula as carteiras Alpha (Rico) e Prime (XP) com cotacoes atuais.
Os dados pessoais (posicoes) ficam em /data, dentro da VPS. Nunca no GitHub.
"""
import io
import json
import os
import threading
import time
from datetime import datetime, date
from zoneinfo import ZoneInfo

from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from pydantic import BaseModel, Field

from config import CARTEIRAS, RENDA_FIXA

app = FastAPI(title="Motor Financeiro")

TOKEN = os.getenv("API_TOKEN", "")
DATA_DIR = os.getenv("DATA_DIR", "/data")
ARQ_POSICOES = os.path.join(DATA_DIR, "posicoes.json")
FUSO = ZoneInfo("America/Sao_Paulo")
_lock = threading.Lock()

# Numero da conta de cada carteira (vem das variaveis de ambiente do Coolify)
CONTAS = {
    os.getenv("CONTA_ALPHA", ""): "alpha",
    os.getenv("CONTA_PRIME", ""): "prime",
}
# Carteira dona dos ativos alugados (a B3 nao informa a conta no aluguel)
CARTEIRA_EMPRESTIMOS = os.getenv("CARTEIRA_EMPRESTIMOS", "alpha")


# ---------------------------------------------------------------- seguranca
def autorizado(x_api_token: str = Header(default="")):
    if not TOKEN or x_api_token != TOKEN:
        raise HTTPException(status_code=401, detail="nao autorizado")


# ---------------------------------------------------------------- arquivos
def agora():
    return datetime.now(FUSO).isoformat(timespec="seconds")


def ler_posicoes():
    if not os.path.exists(ARQ_POSICOES):
        return {"carteiras": {"alpha": {}, "prime": {}}, "tesouro": [],
                "pendentes": [], "aportes_confirmados": [], "atualizado_em": None}
    with open(ARQ_POSICOES, encoding="utf-8") as f:
        return json.load(f)


def salvar_posicoes(dados):
    os.makedirs(DATA_DIR, exist_ok=True)
    tmp = ARQ_POSICOES + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)
    os.replace(tmp, ARQ_POSICOES)


def normalizar_ticker(t: str) -> str:
    t = (t or "").strip().upper()
    # Mercado fracionario: BBAS3F -> BBAS3
    if len(t) >= 6 and t.endswith("F") and t[-2].isdigit():
        t = t[:-1]
    return t


# ---------------------------------------------------------------- cotacoes
_cache = {"quando": 0.0, "dados": {}}
CACHE_SEGUNDOS = 60


def buscar_cotacoes(tickers):
    """Ultimo preco de cada ticker via yfinance, com data/hora da cotacao."""
    tickers = sorted(set(tickers))
    if time.time() - _cache["quando"] < CACHE_SEGUNDOS and all(t in _cache["dados"] for t in tickers):
        return {t: _cache["dados"][t] for t in tickers}

    import yfinance as yf

    simbolos = [t + ".SA" for t in tickers]
    resultado = {}
    try:
        df = yf.download(simbolos, period="5d", interval="15m", group_by="ticker",
                         progress=False, threads=True, auto_adjust=False)
    except Exception:
        df = None
    for t, s in zip(tickers, simbolos):
        try:
            serie = df[s]["Close"].dropna() if len(simbolos) > 1 else df["Close"].dropna()
            if len(serie):
                ts = serie.index[-1]
                ts = ts.tz_convert(FUSO) if ts.tzinfo else ts
                resultado[t] = {"preco": round(float(serie.iloc[-1]), 2),
                                "cotacao_em": ts.isoformat(timespec="minutes")}
        except Exception:
            pass
    _cache["dados"].update(resultado)
    _cache["quando"] = time.time()
    return resultado


# ---------------------------------------------------------------- rotas
@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/teste", dependencies=[Depends(autorizado)])
def teste():
    return {"mensagem": "motor financeiro funcionando"}


@app.get("/carteiras", dependencies=[Depends(autorizado)])
def carteiras():
    dados = ler_posicoes()

    # posicao oficial + pendentes
    pos = {c: {t: dict(p) for t, p in ativos.items()} for c, ativos in dados["carteiras"].items()}
    for pend in dados.get("pendentes", []):
        c, t = pend["carteira"], pend["ticker"]
        p = pos.setdefault(c, {}).setdefault(t, {"tipo": "acao", "qtd": 0, "emprestada": 0})
        p["qtd"] += pend["qtd"]
        p["pendente"] = p.get("pendente", 0) + pend["qtd"]

    todos = [t for ativos in pos.values() for t in ativos]
    cot = buscar_cotacoes(todos)

    saida = {"gerado_em": agora(), "posicao_b3_de": dados.get("data_referencia"),
             "carteiras": {}, "sem_cotacao": []}
    total_rv = 0.0
    for c, ativos in pos.items():
        lista, valor_c = [], 0.0
        for t, p in sorted(ativos.items()):
            q = cot.get(t)
            if not q:
                saida["sem_cotacao"].append(t)
                continue
            valor = round(p["qtd"] * q["preco"], 2)
            valor_c += valor
            lista.append({"ticker": t, "tipo": p.get("tipo"), "qtd": p["qtd"],
                          "emprestada": p.get("emprestada", 0), "pendente": p.get("pendente", 0),
                          "preco": q["preco"], "cotacao_em": q["cotacao_em"], "valor": valor})
        for a in lista:
            a["peso_pct"] = round(100 * a["valor"] / valor_c, 2) if valor_c else 0
        lista.sort(key=lambda a: -a["valor"])
        info = CARTEIRAS.get(c, {"nome": c})
        saida["carteiras"][c] = {**info, "valor": round(valor_c, 2), "qtd_ativos": len(lista), "ativos": lista}
        total_rv += valor_c

    rf = round(sum(x.get("valor_bruto", 0) for x in dados.get("tesouro", [])), 2)
    total = total_rv + rf
    saida["renda_fixa"] = {**RENDA_FIXA, "valor": rf, "titulos": dados.get("tesouro", []),
                           "valor_de": dados.get("data_referencia"),
                           "pct_atual": round(100 * rf / total, 2) if total else 0}
    for c in saida["carteiras"].values():
        c["pct_do_total"] = round(100 * c["valor"] / total, 2) if total else 0
    saida["total"] = round(total, 2)
    return saida


# ---------- aportes (registro rapido, ate a B3 confirmar)
class Aporte(BaseModel):
    carteira: str = Field(..., description="alpha ou prime")
    ticker: str
    qtd: float = Field(..., gt=0)
    preco: float | None = Field(default=None, ge=0)
    data: date | None = None


@app.post("/aportes", dependencies=[Depends(autorizado)])
def registrar_aporte(a: Aporte):
    carteira = a.carteira.strip().lower()
    if carteira not in CARTEIRAS:
        raise HTTPException(400, f"carteira invalida: {a.carteira} (use alpha ou prime)")
    ticker = normalizar_ticker(a.ticker)
    if not ticker:
        raise HTTPException(400, "ticker vazio")
    item = {"data": (a.data or datetime.now(FUSO).date()).isoformat(), "carteira": carteira,
            "ticker": ticker, "qtd": a.qtd, "preco": a.preco, "registrado_em": agora()}
    with _lock:
        dados = ler_posicoes()
        dados.setdefault("pendentes", []).append(item)
        salvar_posicoes(dados)
    return {"ok": True, "registrado": item, "pendentes": len(dados["pendentes"])}


@app.get("/aportes", dependencies=[Depends(autorizado)])
def listar_aportes():
    d = ler_posicoes()
    return {"pendentes": d.get("pendentes", []), "confirmados": d.get("aportes_confirmados", [])}


# ---------- importacao da planilha de posicao da B3
def ler_planilha_b3(conteudo: bytes):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
    nova = {c: {} for c in CARTEIRAS}
    avisos = []

    def linhas(aba):
        if aba not in wb.sheetnames:
            return []
        ws = wb[aba]
        rows = list(ws.iter_rows(values_only=True))
        cab = [str(c or "").strip() for c in rows[0]]
        return [dict(zip(cab, r)) for r in rows[1:] if r and r[0]]

    for aba, tipo in (("Acoes", "acao"), ("Ações", "acao"), ("Fundo de Investimento", "fii")):
        for r in linhas(aba):
            t = normalizar_ticker(str(r.get("Código de Negociação") or ""))
            conta = str(r.get("Conta") or "").strip()
            if not t:
                continue
            c = CONTAS.get(conta)
            if not c:
                avisos.append(f"{t}: conta {conta} nao mapeada (confira CONTA_ALPHA/CONTA_PRIME)")
                continue
            p = nova[c].setdefault(t, {"tipo": tipo, "qtd": 0, "emprestada": 0})
            p["qtd"] += float(r.get("Quantidade") or 0)

    for r in linhas("Empréstimos"):
        t = normalizar_ticker(str(r.get("Produto") or "").split(" ")[0])
        q = float(r.get("Quantidade") or 0)
        if not t or not q:
            continue
        p = nova.setdefault(CARTEIRA_EMPRESTIMOS, {}).setdefault(
            t, {"tipo": "fii" if t.endswith("11") else "acao", "qtd": 0, "emprestada": 0})
        p["qtd"] += q
        p["emprestada"] += q

    tesouro = []
    for r in linhas("Tesouro Direto"):
        if not r.get("Produto") or r.get("Valor Atualizado") in (None, ""):
            continue
        tesouro.append({"titulo": r["Produto"], "vencimento": r.get("Vencimento"),
                        "valor_aplicado": r.get("Valor Aplicado"), "valor_bruto": r.get("Valor bruto"),
                        "valor_liquido": r.get("Valor líquido")})

    for c in nova.values():
        for p in c.values():
            for k in ("qtd", "emprestada"):
                p[k] = int(p[k]) if float(p[k]).is_integer() else p[k]
    return nova, tesouro, avisos


@app.post("/b3/importar", dependencies=[Depends(autorizado)])
async def importar_b3(planilha: UploadFile = File(...)):
    conteudo = await planilha.read()
    try:
        nova, tesouro, avisos = ler_planilha_b3(conteudo)
    except Exception as e:
        raise HTTPException(400, f"nao consegui ler a planilha: {e}")
    if not any(nova.values()):
        raise HTTPException(400, "planilha sem posicoes reconhecidas")

    with _lock:
        dados = ler_posicoes()
        antiga = dados.get("carteiras", {})
        # quanto cada ativo aumentou em relacao a posicao anterior
        aumento = {}
        for c in set(nova) | set(antiga):
            for t in set(nova.get(c, {})) | set(antiga.get(c, {})):
                aumento[(c, t)] = nova.get(c, {}).get(t, {}).get("qtd", 0) - antiga.get(c, {}).get(t, {}).get("qtd", 0)

        primeira_carga = not any(antiga.values())
        confirmados, continuam = [], []
        for pend in sorted(dados.get("pendentes", []), key=lambda x: x["data"]):
            k = (pend["carteira"], pend["ticker"])
            # Na primeira carga nao ha posicao anterior para comparar: pendentes continuam
            if not primeira_carga and aumento.get(k, 0) >= pend["qtd"] - 1e-9:
                aumento[k] -= pend["qtd"]
                confirmados.append({**pend, "confirmado_em": agora()})
            else:
                continuam.append(pend)
        for pend in continuam:
            avisos.append(f"pendente ainda nao aparece na B3: {pend['carteira']} {pend['ticker']} +{pend['qtd']} ({pend['data']})")

        dados.update({"carteiras": nova, "tesouro": tesouro, "pendentes": continuam,
                      "data_referencia": datetime.now(FUSO).date().isoformat(),
                      "arquivo_b3": planilha.filename, "atualizado_em": agora()})
        dados.setdefault("aportes_confirmados", []).extend(confirmados)
        salvar_posicoes(dados)

    return {"ok": True, "arquivo": planilha.filename,
            "ativos": {c: len(v) for c, v in nova.items()}, "titulos_tesouro": len(tesouro),
            "pendentes_confirmados": len(confirmados), "pendentes_restantes": len(continuam),
            "avisos": avisos}

# ============================================================ REGRA DE APORTE
# (bloco adicionado - cole no FINAL do main.py)
import aporte as regra_aporte


class PedidoAporte(BaseModel):
    carteira: str = Field(..., description="alpha ou prime")
    valor: float = Field(..., gt=0, description="quanto voce quer aportar (R$)")
    max_ativos: int = Field(default=5, ge=1, le=20, description="em quantos ativos dividir")


@app.post("/aporte/sugestao", dependencies=[Depends(autorizado)])
def sugestao_aporte(p: PedidoAporte):
    c = p.carteira.strip().lower()
    if c not in CARTEIRAS:
        raise HTTPException(400, "carteira invalida (use alpha ou prime)")
    visao = carteiras()
    ativos = visao["carteiras"].get(c, {}).get("ativos", [])
    if not ativos:
        raise HTTPException(400, "carteira sem ativos com cotacao")
    inds = regra_aporte.indicadores([a["ticker"] for a in ativos], DATA_DIR)
    r = regra_aporte.sugerir(c, p.valor, ativos, inds, p.max_ativos)
    r["gerado_em"] = agora()
    r["objetivo"] = CARTEIRAS[c]["objetivo"]
    return r


@app.get("/indicadores/{carteira}", dependencies=[Depends(autorizado)])
def ver_indicadores(carteira: str):
    c = carteira.strip().lower()
    if c not in CARTEIRAS:
        raise HTTPException(400, "carteira invalida (use alpha ou prime)")
    tickers = list(ler_posicoes()["carteiras"].get(c, {}).keys())
    inds = regra_aporte.indicadores(tickers, DATA_DIR)
    nt = regra_aporte.notas(c, inds)
    return {"carteira": c, "gerado_em": agora(),
            "ativos": sorted([{**inds[t], **nt[t]} for t in tickers], key=lambda x: -x["nota"])}


@app.delete("/aportes/pendente", dependencies=[Depends(autorizado)])
def apagar_pendente(registrado_em: str):
    """Apaga um pendente pelo campo registrado_em (veja em GET /aportes)."""
    with _lock:
        dados = ler_posicoes()
        antes = len(dados.get("pendentes", []))
        dados["pendentes"] = [x for x in dados.get("pendentes", []) if x.get("registrado_em") != registrado_em]
        if len(dados["pendentes"]) == antes:
            raise HTTPException(404, "pendente nao encontrado")
        salvar_posicoes(dados)
    return {"ok": True, "pendentes": len(dados["pendentes"])}
